#!/usr/bin/env python3
"""Build the online Tiverton Town results centre from public web sources.

Designed for GitHub Actions/Linux. It fetches:
- Football Web Pages match grid + monthly fixture/result pages
- Tivvy Archive league table
Then writes public/data.json, public/data.js, matrix CSV, and formatted XLSX.
"""
from __future__ import annotations

import csv
import html
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
PUBLIC = ROOT / "public"
PUBLIC.mkdir(exist_ok=True)

SOURCE_URL = "https://www.footballwebpages.co.uk/southern-football-league-division-one-south/match-grid"
FIXTURES_BASE_URL = "https://www.footballwebpages.co.uk/southern-football-league-division-one-south/fixtures-results"
LEAGUE_SOURCE_URL = "https://www.tivvyarchive.co.uk/season.php?year=2026"

CODES = ["BAR","BID","BIS","BRI","DOR","EXM","FAL","HAR","HUN","LAR","MEL","PAU","POR","SHA","SLI","SPO","SWI","TIV","WES","WEY","WIL","WOR"]
NAMES = {
    "BAR":"Barnstaple Town", "BID":"Bideford", "BIS":"Bishops Cleeve", "BRI":"Bristol Manor Farm",
    "DOR":"Dorchester", "EXM":"Exmouth Town", "FAL":"Falmouth Town", "HAR":"Hartpury", "HUN":"Hungerford",
    "LAR":"Larkhall Athletic", "MEL":"Melksham Town", "PAU":"Paulton Rovers", "POR":"Portland Utd",
    "SHA":"Shaftesbury", "SLI":"Slimbridge AFC", "SPO":"Sporting Club Inkberrow", "SWI":"Swindon SM",
    "TIV":"Tiverton Town", "WES":"Westbury", "WEY":"Weymouth FC", "WIL":"Willand Rovers", "WOR":"Worcester Raiders"
}
ALIASES = {
    "Barnstaple Town":"Barnstaple Town", "Bideford":"Bideford", "Bishops Cleeve":"Bishops Cleeve",
    "Bristol Manor Farm":"Bristol Manor Farm", "Dorchester":"Dorchester", "Dorchester Town":"Dorchester",
    "Exmouth Town":"Exmouth Town", "Falmouth Town":"Falmouth Town", "Hartpury":"Hartpury",
    "Hungerford":"Hungerford", "Hungerford Town":"Hungerford", "Larkhall Athletic":"Larkhall Athletic",
    "Melksham Town":"Melksham Town", "Paulton Rovers":"Paulton Rovers", "Portland Utd":"Portland Utd",
    "Portland United":"Portland Utd", "Shaftesbury":"Shaftesbury", "Slimbridge AFC":"Slimbridge AFC",
    "Slimbridge":"Slimbridge AFC", "Sporting Club Inkberrow":"Sporting Club Inkberrow", "Swindon SM":"Swindon SM",
    "Swindon Supermarine":"Swindon SM", "Tiverton Town":"Tiverton Town", "Westbury":"Westbury",
    "Westbury United":"Westbury", "Weymouth FC":"Weymouth FC", "Weymouth":"Weymouth FC",
    "Willand Rovers":"Willand Rovers", "Worcester Raiders":"Worcester Raiders"
}
ALIASES_BY_CANON = {}
for alias, canon in ALIASES.items():
    ALIASES_BY_CANON.setdefault(canon, []).append(alias)

MONTHS = ["august","september","october","november","december","january","february","march","april"]
DATE_RE = re.compile(r"(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\s+\d{1,2}(?:st|nd|rd|th)\s+(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+20\d{2}", re.I)
SCORE_RE = re.compile(r"^(\d+)\s*-\s*(\d+)$")
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; TivertonTownResults/1.0)"}


def get(url: str, timeout: int = 30) -> str:
    r = requests.get(url, headers=HEADERS, timeout=timeout)
    r.raise_for_status()
    return r.text


def clean_text(value: str) -> str:
    if value is None:
        return ""
    value = html.unescape(value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def parse_match_grid(page: str) -> list[dict]:
    soup = BeautifulSoup(page, "html.parser")
    target = None
    for table in soup.find_all("table"):
        text = table.get_text(" ", strip=True)
        if "BAR" in text and "BID" in text and "WOR" in text:
            target = table
            break
    if target is None:
        raise RuntimeError("Could not find the Southern League Division One South match grid")

    rows = []
    for tr in target.find_all("tr"):
        cells = [clean_text(c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"])]
        if len(cells) < len(CODES) + 1:
            continue
        home_code = cells[0].upper()
        if home_code not in CODES:
            continue
        for idx, away_code in enumerate(CODES):
            cell = cells[idx + 1]
            m = SCORE_RE.match(cell)
            if m:
                rows.append({"date":"", "home":NAMES[home_code], "away":NAMES[away_code], "score":f"{m.group(1)}-{m.group(2)}"})

    dedup = {(r["home"], r["away"], r["score"]): r for r in rows}
    result = list(dedup.values())
    if not result:
        raise RuntimeError("Match grid found, but no completed scores were present")
    return result


def parse_date(text: str) -> str:
    text = re.sub(r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\s+", "", text, flags=re.I)
    text = re.sub(r"(\d{1,2})(st|nd|rd|th)", r"\1", text, flags=re.I)
    return datetime.strptime(text.strip(), "%d %B %Y").strftime("%Y-%m-%d")


def fixture_in_block(block: str, fixture: dict) -> bool:
    m = SCORE_RE.match(fixture["score"])
    if not m:
        return False
    hs, as_ = map(re.escape, m.groups())
    for ha in ALIASES_BY_CANON.get(fixture["home"], [fixture["home"]]):
        for aa in ALIASES_BY_CANON.get(fixture["away"], [fixture["away"]]):
            pat = (r"(?is)(?<![\w])" + re.escape(ha) + r"(?![\w]).{0,90}?(?<!\d)" + hs +
                   r"(?!\d).{0,35}?(?<!\d)" + as_ + r"(?!\d).{0,90}?(?<![\w])" + re.escape(aa) + r"(?![\w])")
            if re.search(pat, block):
                return True
    return False


def add_dates(results: list[dict]) -> None:
    for month in MONTHS:
        try:
            page = get(f"{FIXTURES_BASE_URL}/{month}", 25)
        except Exception as exc:
            print(f"{month}: unavailable ({exc})")
            continue
        soup = BeautifulSoup(page, "html.parser")
        text = soup.get_text("\n", strip=True)
        matches = list(DATE_RE.finditer(text))
        if not matches:
            print(f"{month}: no date headings recognised")
            continue
        for i, dm in enumerate(matches):
            try:
                iso = parse_date(dm.group(0))
            except ValueError:
                continue
            start = dm.end()
            end = matches[i+1].start() if i+1 < len(matches) else len(text)
            block = text[start:end]
            for fixture in results:
                if not fixture["date"] and fixture_in_block(block, fixture):
                    fixture["date"] = iso


def parse_league_table(page: str) -> list[dict]:
    soup = BeautifulSoup(page, "html.parser")
    heading = None
    for h in soup.find_all(re.compile(r"^h[1-6]$")):
        if clean_text(h.get_text()).lower() == "league table":
            heading = h
            break
    table = heading.find_next("table") if heading else None
    if table is None:
        return []
    out = []
    for tr in table.find_all("tr"):
        cells = [clean_text(c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"])]
        if len(cells) < 5 or not cells[0].isdigit() or not cells[2].isdigit() or not re.fullmatch(r"[+-]?\d+", cells[3]) or not cells[4].isdigit():
            continue
        out.append({"pos":int(cells[0]), "team":cells[1], "p":int(cells[2]), "gd":int(cells[3]), "pts":int(cells[4])})
    return sorted(out, key=lambda x: x["pos"])


def load_previous() -> dict:
    p = PUBLIC / "data.json"
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8-sig"))
    except Exception:
        return {}


def build_csv(results: list[dict], league: list[dict]) -> None:
    score_lookup = {}
    code_by_name = {v:k for k,v in NAMES.items()}
    for x in results:
        hc, ac = code_by_name.get(x["home"]), code_by_name.get(x["away"])
        if hc and ac:
            score_lookup[(hc, ac)] = x["score"]
    out = PUBLIC / "SFL_Division_One_South_Matrix.csv"
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["A/H"] + CODES)
        for away in CODES:
            row = [away]
            for home in CODES:
                row.append("" if home == away else score_lookup.get((home, away), ""))
            w.writerow(row)
        w.writerow([])
        w.writerow(["League table"])
        w.writerow(["Pos","Team","P","GD","Pts"])
        for x in league:
            w.writerow([x["pos"],x["team"],x["p"],x["gd"],x["pts"]])


def build_xlsx(results: list[dict], league: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results Matrix"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"

    yellow = "F2C400"; axis_corner = "FF8C42"; tiv_blue = "7FC8F8"; white = "FFFFFF"; grid = "333333"
    auto_p = "9BD18B"; promo_po = "9DC3E6"; releg_po = "F4B183"; releg = "F4A6A6"
    thin = Side(style="thin", color=grid)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    score_lookup = {}
    code_by_name = {v:k for k,v in NAMES.items()}
    for x in results:
        hc, ac = code_by_name.get(x["home"]), code_by_name.get(x["away"])
        if hc and ac:
            score_lookup[(hc, ac)] = x["score"]

    ws.cell(1,1,"A/H")
    for i, code in enumerate(CODES, start=2):
        ws.cell(1,i,code); ws.cell(i,1,code)

    max_r = len(CODES) + 1; max_c = len(CODES) + 1
    for row in ws.iter_rows(min_row=1,max_row=max_r,min_col=1,max_col=max_c):
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = center
            cell.border = border
            cell.fill = PatternFill("solid", fgColor=white)
            cell.number_format = "@"

    for c in range(1,max_c+1):
        ws.cell(1,c).fill = PatternFill("solid", fgColor=yellow); ws.cell(1,c).font = Font(name="Arial", size=11, bold=True)
    for r in range(1,max_r+1):
        ws.cell(r,1).fill = PatternFill("solid", fgColor=yellow); ws.cell(r,1).font = Font(name="Arial", size=11, bold=True)
    ws.cell(1,1).fill = PatternFill("solid", fgColor=axis_corner)

    tiv_idx = CODES.index("TIV") + 2
    ws.cell(1,tiv_idx).fill = PatternFill("solid", fgColor=tiv_blue)
    ws.cell(tiv_idx,1).fill = PatternFill("solid", fgColor=tiv_blue)

    for ai, away in enumerate(CODES, start=2):
        for hi, home in enumerate(CODES, start=2):
            cell = ws.cell(ai,hi)
            if home == away:
                cell.fill = PatternFill("solid", fgColor=yellow)
            elif (home,away) in score_lookup:
                cell.value = score_lookup[(home,away)]
                cell.font = Font(name="Arial", size=11, bold=True)

    for col in range(1,max_c+1):
        ws.column_dimensions[get_column_letter(col)].width = 7.5
    for r in range(1,max_r+1):
        ws.row_dimensions[r].height = 25

    title_row = max_r + 3
    head_row = title_row + 1
    data_row = head_row + 1
    ws.merge_cells(start_row=title_row,start_column=1,end_row=title_row,end_column=7)
    c = ws.cell(title_row,1,"League table")
    c.fill=PatternFill("solid",fgColor=yellow); c.font=Font(name="Arial",size=14,bold=True); c.alignment=center
    ws.cell(head_row,1,"Pos"); ws.merge_cells(start_row=head_row,start_column=2,end_row=head_row,end_column=4); ws.cell(head_row,2,"Team")
    ws.cell(head_row,5,"P"); ws.cell(head_row,6,"GD"); ws.cell(head_row,7,"Pts")
    for col in range(1,8):
        cell=ws.cell(head_row,col); cell.fill=PatternFill("solid",fgColor=yellow); cell.font=Font(name="Arial",size=11,bold=True); cell.alignment=center; cell.border=border

    club_count = len(league)
    for offset, x in enumerate(league):
        r = data_row + offset
        ws.cell(r,1,str(x["pos"])); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); ws.cell(r,2,x["team"])
        ws.cell(r,5,str(x["p"])); ws.cell(r,6,str(x["gd"])); ws.cell(r,7,str(x["pts"]))
        pos=x["pos"]; fill=None
        if pos==1: fill=auto_p
        elif 2<=pos<=5: fill=promo_po
        elif club_count>=4 and club_count-3<=pos<=club_count-2: fill=releg_po
        elif club_count>=2 and pos>=club_count-1: fill=releg
        for col in range(1,8):
            cell=ws.cell(r,col); cell.font=Font(name="Arial",size=11,bold=(x["team"]=="Tiverton Town")); cell.alignment=center; cell.border=border
            if fill: cell.fill=PatternFill("solid",fgColor=fill)
        ws.cell(r,2).alignment=Alignment(horizontal="left",vertical="center")
        if x["team"]=="Tiverton Town": ws.cell(r,2).fill=PatternFill("solid",fgColor=tiv_blue)
        ws.row_dimensions[r].height=23

    wb.save(PUBLIC / "SFL_Division_One_South_Matrix.xlsx")


def main() -> int:
    prev = load_previous()
    try:
        results = parse_match_grid(get(SOURCE_URL))
        add_dates(results)
    except Exception as exc:
        print(f"Results update failed: {exc}", file=sys.stderr)
        results = prev.get("results", [])
        if not results:
            return 1

    try:
        league = parse_league_table(get(LEAGUE_SOURCE_URL))
    except Exception as exc:
        print(f"League table update failed: {exc}", file=sys.stderr)
        league = []
    if not league:
        league = prev.get("leagueTable", [])

    results.sort(key=lambda x: (x.get("date") or "0000-00-00", x.get("home", ""), x.get("away", "")), reverse=True)
    results = sorted(results, key=lambda x: x.get("away", ""))
    results = sorted(results, key=lambda x: x.get("home", ""))
    results = sorted(results, key=lambda x: x.get("date") or "0000-00-00", reverse=True)

    data = {
        "teams":[NAMES[c] for c in CODES],
        "codes":{NAMES[c]:c for c in CODES},
        "results":results,
        "leagueTable":league,
        "updated":datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        "source":SOURCE_URL,
        "leagueSource":LEAGUE_SOURCE_URL,
    }
    raw = json.dumps(data, ensure_ascii=False, separators=(",",":"))
    (PUBLIC/"data.json").write_text(raw, encoding="utf-8")
    (PUBLIC/"data.js").write_text("window.SFL_DATA = " + raw + ";\n", encoding="utf-8")
    build_csv(results, league)
    build_xlsx(results, league)
    dated = sum(bool(x.get("date")) for x in results)
    print(f"Success: {len(results)} results ({dated} dated), {len(league)} league rows")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
